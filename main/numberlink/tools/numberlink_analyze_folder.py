#!/usr/bin/env python3
"""
numberlink_analyze_folder.py — フォルダ内のパズルを分析してExcel出力
====================================================================

使用例:
  python numberlink_analyze_folder.py ./puzzles -o analysis.xlsx
  python numberlink_analyze_folder.py ../src/puzzles/data  # デフォルト: ../analysis/difficulty_analysis.xlsx
"""

import os
import sys
import argparse
import json
from pathlib import Path

_dir = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _dir)

# 出力先はカレントではなく ../analysis/ 固定（分析データはリポジトリに上げない）
_DEFAULT_OUTPUT = os.path.normpath(
    os.path.join(_dir, "..", "analysis", "difficulty_analysis.xlsx"))

from numberlink_difficulty import analyze_puzzle
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# 出力する指標とその表示名
METRICS = [
    ("file", "ファイル名"),
    ("num_pairs", "ペア数"),
    ("usable_cells", "使用可能セル"),
    ("manhattan_min", "Manhattan最小"),
    ("manhattan_max", "Manhattan最大"),
    ("manhattan_avg", "Manhattan平均"),
    ("endpoints_on_edge_pct", "辺上端点(%)"),
    ("endpoints_on_corner", "角の端点数"),
    ("min_inter_endpoint_dist", "異ペア端点最小距離"),
    ("bottleneck_cells", "ボトルネック数"),
    ("bottleneck_pct", "ボトルネック(%)"),
    ("path_len_min", "パス長最小"),
    ("path_len_max", "パス長最大"),
    ("path_len_avg", "パス長平均"),
    ("path_len_std", "パス長標準偏差"),
    ("detour_ratio_avg", "迂回度平均"),
    ("detour_ratio_max", "迂回度最大"),
    ("inter_path_contacts", "パス間接触数"),
    ("total_turns", "曲がり回数"),
    ("turns_per_cell", "曲がり/セル"),
    ("sat_decisions", "SAT決定回数"),
    ("sat_conflicts", "SAT競合回数"),
    ("sat_propagations", "SAT伝播回数"),
    ("solve_time_ms", "SAT解答時間(ms)"),
    ("propagation_determined", "理詰め確定数"),
    ("propagation_total_empty", "空きセル総数"),
    ("propagation_ratio", "理詰め確定率(%)"),
    ("propagation_reduced", "候補減少セル数"),
    ("propagation_avg_remaining", "未確定平均候補数"),
]


def find_puzzles(folder: str) -> list:
    """フォルダ内のJSONパズルファイルを検索"""
    puzzles = []
    folder_path = Path(folder)
    
    for json_file in sorted(folder_path.glob("*.json")):
        try:
            with open(json_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
            if "size" in data and "numbers" in data:
                puzzles.append(str(json_file))
        except (json.JSONDecodeError, KeyError):
            continue
    
    return puzzles


def get_puzzle_size(filepath: str) -> str:
    """パズルのサイズを取得"""
    with open(filepath, 'r', encoding='utf-8') as f:
        data = json.load(f)
    rows, cols = data["size"]
    return f"{rows}x{cols}"


def create_xlsx(results_by_size: dict, output_path: str):
    """サイズごとにシートを作成してExcel出力"""
    wb = Workbook()
    wb.remove(wb.active)
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    number_fill_light = PatternFill("solid", fgColor="D6DCE4")
    border = Side(style="thin", color="B4B4B4")
    cell_border = Border(left=border, right=border, top=border, bottom=border)
    
    # 重要指標のハイライト列
    highlight_cols = {"SAT決定回数", "SAT競合回数", "理詰め確定率(%)"}
    highlight_fill = PatternFill("solid", fgColor="FFF2CC")
    
    for size in sorted(results_by_size.keys(), key=lambda s: int(s.split('x')[0])):
        results = results_by_size[size]
        ws = wb.create_sheet(title=size)
        
        # ヘッダー行
        for col_idx, (key, label) in enumerate(METRICS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = cell_border
        
        # データ行
        for row_idx, metrics in enumerate(results, start=2):
            for col_idx, (key, label) in enumerate(METRICS, start=1):
                value = metrics.get(key, "")
                if isinstance(value, float):
                    value = round(value, 2)
                
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = cell_border
                
                # 数値は右寄せ
                if isinstance(value, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
                
                # 偶数行に背景色
                if row_idx % 2 == 0:
                    cell.fill = number_fill_light
                
                # 重要指標をハイライト
                if label in highlight_cols:
                    cell.fill = highlight_fill
        
        # 列幅調整
        for col_idx, (key, label) in enumerate(METRICS, start=1):
            col_letter = get_column_letter(col_idx)
            if key == "file":
                ws.column_dimensions[col_letter].width = 25
            elif len(label) > 10:
                ws.column_dimensions[col_letter].width = 14
            else:
                ws.column_dimensions[col_letter].width = 10
        
        # フィルター設定
        ws.auto_filter.ref = f"A1:{get_column_letter(len(METRICS))}{len(results)+1}"
        
        # 先頭行を固定
        ws.freeze_panes = "B2"
    
    # サマリーシートを先頭に追加
    summary = wb.create_sheet(title="サマリー", index=0)
    summary["A1"] = "サイズ"
    summary["B1"] = "問題数"
    summary["C1"] = "SAT決定(平均)"
    summary["D1"] = "SAT競合(平均)"
    summary["E1"] = "理詰め確定率(平均)"
    
    for col in range(1, 6):
        cell = summary.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = cell_border
    
    row = 2
    for size in sorted(results_by_size.keys(), key=lambda s: int(s.split('x')[0])):
        results = results_by_size[size]
        summary.cell(row=row, column=1, value=size).border = cell_border
        summary.cell(row=row, column=2, value=len(results)).border = cell_border
        
        avg_decisions = sum(r.get("sat_decisions", 0) for r in results) / len(results)
        avg_conflicts = sum(r.get("sat_conflicts", 0) for r in results) / len(results)
        avg_ratio = sum(r.get("propagation_ratio", 0) for r in results) / len(results)
        
        summary.cell(row=row, column=3, value=round(avg_decisions, 1)).border = cell_border
        summary.cell(row=row, column=4, value=round(avg_conflicts, 1)).border = cell_border
        summary.cell(row=row, column=5, value=round(avg_ratio, 1)).border = cell_border
        row += 1
    
    for col, width in [(1, 10), (2, 10), (3, 15), (4, 15), (5, 18)]:
        summary.column_dimensions[get_column_letter(col)].width = width
    
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(
        description="フォルダ内のナンバーリンクパズルを分析してExcel出力")
    parser.add_argument("folder", help="パズルJSONが入ったフォルダ")
    parser.add_argument("-o", "--output", default=_DEFAULT_OUTPUT,
                        help=f"出力ファイル名 (デフォルト: {_DEFAULT_OUTPUT})")
    
    args = parser.parse_args()
    
    if not os.path.isdir(args.folder):
        print(f"エラー: フォルダが見つかりません: {args.folder}")
        sys.exit(1)
    
    puzzles = find_puzzles(args.folder)
    if not puzzles:
        print(f"エラー: パズルファイルが見つかりません: {args.folder}")
        sys.exit(1)
    
    print(f"フォルダ: {args.folder}")
    print(f"パズル数: {len(puzzles)}")
    print()
    
    # サイズごとに分類しながら分析
    results_by_size = {}
    
    for i, puzzle_path in enumerate(puzzles, start=1):
        size = get_puzzle_size(puzzle_path)
        filename = os.path.basename(puzzle_path)
        print(f"  [{i}/{len(puzzles)}] {filename} ({size})...", end=" ", flush=True)
        
        try:
            metrics = analyze_puzzle(puzzle_path)
            if size not in results_by_size:
                results_by_size[size] = []
            results_by_size[size].append(metrics)
            print("OK")
        except Exception as e:
            print(f"エラー: {e}")
    
    print()
    
    # Excel出力
    output_path = args.output
    if not output_path.endswith(".xlsx"):
        output_path += ".xlsx"

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    create_xlsx(results_by_size, output_path)
    print(f"出力: {output_path}")
    
    # サマリー表示
    print()
    print("=== サマリー ===")
    for size in sorted(results_by_size.keys(), key=lambda s: int(s.split('x')[0])):
        results = results_by_size[size]
        avg_decisions = sum(r.get("sat_decisions", 0) for r in results) / len(results)
        avg_ratio = sum(r.get("propagation_ratio", 0) for r in results) / len(results)
        print(f"  {size}: {len(results)}問, "
              f"SAT決定平均={avg_decisions:.1f}, "
              f"理詰め確定率={avg_ratio:.1f}%")


if __name__ == "__main__":
    main()
